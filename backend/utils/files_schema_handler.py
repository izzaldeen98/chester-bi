"""
files_schema_handler.py

Infers the column schema (name + type) of an uploaded data file — csv, json,
parquet, xlsx — by reading only a small, bounded snippet of it rather than
the whole payload, so a "preview schema" API stays fast even for multi-GB
analytical files sitting in S3 or on local disk.

Each format has its own reader (strategy pattern, one class per extension)
because *where* the useful bytes live differs by format:
  - csv/json  -> a header + a few sample rows always live at the START.
  - parquet   -> the schema lives in a footer at the END of the file.
  - xlsx      -> a zip container whose sheet data can live anywhere inside
                 it; there is no cheap byte-range trick, so it's read in full
                 (see XlsxSchemaReader for the tradeoff).
"""

from __future__ import annotations

import csv
import io
import json
import re
from abc import ABC, abstractmethod
from dataclasses import dataclass
from typing import Dict, List, Optional

from utils.config_files import storage

# ── Result shape ─────────────────────────────────────────────────────────

# Generic, storage-agnostic type vocabulary — deliberately not tied to any
# single database's or Malloy's type names, since this schema exists before
# a file is connected/imported to anything.
FileColumnType = str  # "string" | "integer" | "float" | "boolean" | "date" | "datetime" | "null"


@dataclass
class FileColumn:
    name: str
    type: FileColumnType

    def to_dict(self) -> dict:
        return {"name": self.name, "type": self.type}


@dataclass
class FileSchema:
    format: str
    columns: List[FileColumn]
    sample_row_count: int

    def to_dict(self) -> dict:
        return {
            "format": self.format,
            "sample_row_count": self.sample_row_count,
            "columns": [column.to_dict() for column in self.columns],
        }


# ── Shared scalar type inference (csv / json / xlsx all reuse this) ────────

_BOOL_VALUES = {"true", "false"}
_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}([T ]\d{2}:\d{2}(:\d{2})?)?$")


def _infer_scalar_type(value: object) -> FileColumnType:
    if value is None:
        return "null"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, int):
        return "integer"
    if isinstance(value, float):
        return "float"
    if not isinstance(value, str):
        return "string"

    text = value.strip()
    if text == "":
        return "null"
    if text.lower() in _BOOL_VALUES:
        return "boolean"
    try:
        int(text)
        return "integer"
    except ValueError:
        pass
    try:
        float(text)
        return "float"
    except ValueError:
        pass
    if _DATE_RE.match(text):
        return "datetime" if (":" in text) else "date"
    return "string"


def _widen_type(current: Optional[FileColumnType], observed: FileColumnType) -> FileColumnType:
    """Merge one more observed value's type into a column's running inferred type."""
    if current is None or current == observed:
        return observed
    if current == "null":
        return observed
    if observed == "null":
        return current
    numeric = {"integer", "float"}
    if current in numeric and observed in numeric:
        return "float"
    return "string"  # any other mismatch -> the safe, always-valid fallback


# ── Reader strategy interface ───────────────────────────────────────────────

class BaseFileSchemaReader(ABC):
    """
    One reader per file extension. `read_schema` only ever sees the bounded
    snippet `_fetch_snippet` decided to fetch for this reader — never the
    full file — so readers must tolerate a truncated/partial payload.
    """

    @abstractmethod
    def read_schema(self, snippet: bytes, sample_rows: int) -> List[FileColumn]:
        ...


class CsvSchemaReader(BaseFileSchemaReader):
    def read_schema(self, snippet: bytes, sample_rows: int) -> List[FileColumn]:
        text = snippet.decode("utf-8", errors="replace")

        # The snippet is a byte prefix, not a whole-line boundary — the last
        # line is very likely cut mid-row, so drop it unless it's all we have.
        lines = text.splitlines()
        if len(lines) > 1:
            lines = lines[:-1]
        sample_text = "\n".join(lines)
        if not sample_text.strip():
            raise ValueError("Not enough data to infer a CSV schema from this snippet")

        try:
            dialect = csv.Sniffer().sniff(sample_text)
        except csv.Error:
            dialect = csv.excel

        reader = csv.DictReader(io.StringIO(sample_text), dialect=dialect)
        if not reader.fieldnames:
            raise ValueError("CSV file has no header row")

        column_types: Dict[str, Optional[FileColumnType]] = {name: None for name in reader.fieldnames}
        rows_seen = 0
        for row in reader:
            if rows_seen >= sample_rows:
                break
            for name in reader.fieldnames:
                column_types[name] = _widen_type(column_types[name], _infer_scalar_type(row.get(name)))
            rows_seen += 1

        return [FileColumn(name=name, type=column_types[name] or "string") for name in reader.fieldnames]


class JsonSchemaReader(BaseFileSchemaReader):
    def read_schema(self, snippet: bytes, sample_rows: int) -> List[FileColumn]:
        text = snippet.decode("utf-8", errors="replace")
        records = self._extract_records(text, sample_rows)
        if not records:
            raise ValueError("Not enough data to infer a JSON schema from this snippet")

        column_types: Dict[str, Optional[FileColumnType]] = {}
        for record in records[:sample_rows]:
            if not isinstance(record, dict):
                continue
            for key, value in record.items():
                column_types[key] = _widen_type(column_types.get(key), _infer_scalar_type(value))

        if not column_types:
            raise ValueError("JSON sample did not contain any object records")
        return [FileColumn(name=name, type=col_type or "string") for name, col_type in column_types.items()]

    @staticmethod
    def _extract_records(text: str, limit: int) -> List[dict]:
        # JSON Lines (one object per line) is the only shape that's actually
        # snippet-friendly — each line parses independently, so a truncated
        # final line can just be skipped.
        records: List[dict] = []
        for line in text.splitlines():
            if not line.strip():
                continue
            try:
                records.append(json.loads(line))
            except json.JSONDecodeError:
                continue  # most likely the snippet's truncated last line
            if len(records) >= limit:
                break
        if records:
            return records

        # Best-effort fallback for a top-level JSON array: only succeeds if
        # the snippet happened to end on a complete element.
        try:
            parsed = json.loads(text)
        except json.JSONDecodeError:
            end = text.rfind("},")
            if end == -1:
                return []
            try:
                parsed = json.loads(text[: end + 1] + "]")
            except json.JSONDecodeError:
                return []

        if isinstance(parsed, list):
            return parsed[:limit]
        if isinstance(parsed, dict):
            return [parsed]
        return []


class ParquetSchemaReader(BaseFileSchemaReader):
    """Parquet's schema lives in a footer at the END of the file — see _fetch_snippet."""

    def read_schema(self, snippet: bytes, sample_rows: int) -> List[FileColumn]:
        try:
            import pyarrow.parquet as pq
        except ImportError as exc:
            raise RuntimeError(
                "Parquet schema inference requires the 'pyarrow' package"
            ) from exc

        parquet_file = pq.ParquetFile(io.BytesIO(snippet))
        arrow_schema = parquet_file.schema_arrow
        return [
            FileColumn(name=name, type=_arrow_type_to_generic(str(arrow_type)))
            for name, arrow_type in zip(arrow_schema.names, arrow_schema.types)
        ]


def _arrow_type_to_generic(arrow_type: str) -> FileColumnType:
    t = arrow_type.lower()
    if "bool" in t:
        return "boolean"
    if "timestamp" in t:
        return "datetime"
    if "date" in t:
        return "date"
    if "int" in t:
        return "integer"
    if "float" in t or "double" in t or "decimal" in t:
        return "float"
    return "string"


class XlsxSchemaReader(BaseFileSchemaReader):
    """
    ponytail: no cheap byte-range trick exists for xlsx — it's a zip
    container whose sheet data can live anywhere inside it, so this reads
    the whole object (see _fetch_snippet). Ceiling: a full download per
    schema preview. Upgrade path: a seekable range-fetching S3 file wrapper
    (e.g. smart_open) if large xlsx previews become a real cost.
    """

    def read_schema(self, snippet: bytes, sample_rows: int) -> List[FileColumn]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError(
                "XLSX schema inference requires the 'openpyxl' package"
            ) from exc

        workbook = load_workbook(io.BytesIO(snippet), read_only=True, data_only=True)
        try:
            sheet = workbook.active
            rows = sheet.iter_rows(max_row=sample_rows + 1, values_only=True)

            try:
                header = next(rows)
            except StopIteration:
                raise ValueError("XLSX sheet is empty")

            column_names = [
                str(cell) if cell is not None else f"column_{i}" for i, cell in enumerate(header)
            ]
            column_types: Dict[str, Optional[FileColumnType]] = {name: None for name in column_names}

            for row in rows:
                for name, cell in zip(column_names, row):
                    column_types[name] = _widen_type(column_types[name], _infer_scalar_type(cell))

            return [FileColumn(name=name, type=column_types[name] or "string") for name in column_names]
        finally:
            workbook.close()


# ── Dispatch ─────────────────────────────────────────────────────────────

_READERS: Dict[str, BaseFileSchemaReader] = {
    "csv": CsvSchemaReader(),
    "json": JsonSchemaReader(),
    "parquet": ParquetSchemaReader(),
    "xlsx": XlsxSchemaReader(),
}

# How much of the file each format needs fetched before parsing.
_PREFIX_SNIPPET_BYTES = 65_536       # csv/json: header + sample rows live at the start
_PARQUET_FOOTER_BYTES = 1_048_576    # comfortably covers real-world Parquet footers


async def get_file_schema(
    path: str,
    file_name: str,
    extension: str,
    sample_rows: int = 50,
) -> FileSchema:
    """
    Infers column names + types for a stored file without downloading it in
    full (csv/json/parquet — see module docstring for the xlsx exception).

    `path` / `file_name` match the storage layer's addressing (utils.config_files).
    `sample_rows` bounds how many data rows are used to widen each column's
    inferred type — more rows means safer inference at the cost of needing
    more snippet bytes for csv/json.
    """
    ext = extension.lower().lstrip(".")
    reader = _READERS.get(ext)
    if reader is None:
        raise ValueError(f"Unsupported file extension for schema inference: '{extension}'")

    snippet = await _fetch_snippet(path, file_name, ext)
    if not snippet:
        raise ValueError(f"File not found or empty: {path}/{file_name}")

    columns = reader.read_schema(snippet, sample_rows)
    return FileSchema(format=ext, columns=columns, sample_row_count=sample_rows)


async def _fetch_snippet(path: str, file_name: str, ext: str) -> Optional[bytes]:
    if ext in ("csv", "json"):
        data = await storage.get_file_range(path, file_name, start=0, length=_PREFIX_SNIPPET_BYTES)
    elif ext == "parquet":
        data = await storage.get_file_range(path, file_name, start=-_PARQUET_FOOTER_BYTES)
    else:
        # xlsx (and any future whole-file-only format)
        data = await storage.get_file(path, file_name)

    return data.getvalue() if data is not None else None
